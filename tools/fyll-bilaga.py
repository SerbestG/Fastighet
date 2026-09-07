"""
Fyller i svarskolumnen i Bilaga 3 och skriver om kravuppfyllnaden.

Båda utdata kommer ur samma källa (kravsvar.py), så att bilagan och
dokumentationen aldrig kan säga olika saker.

Körs med:
    python3 tools/fyll-bilaga.py <indata.xlsx> <utdata.xlsx>
"""

from __future__ import annotations

import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from kravsvar import SVAR, UPPFYLLD, RUTIN, KONFIGURATION, AVTAL, LEVERANS  # noqa: E402

SHEETS = ['A. Icke-funktionella krav', 'B. Funktionella krav', 'C. Informationssäkerhetskrav']

# Kolumner i bilagan.
COL_ID, COL_TEXT, COL_TYPE, COL_ANSWER, COL_COMMENT = 2, 3, 4, 5, 6

ORDNING = [UPPFYLLD, KONFIGURATION, AVTAL, LEVERANS, RUTIN]

FORKLARING = {
    UPPFYLLD: 'Finns i den levererade produkten och går att kontrollera.',
    KONFIGURATION: 'Byggt och klart, men behöver adress och autentiseringsuppgifter.',
    AVTAL: 'Kräver avtal eller certifikat som ännu inte finns.',
    LEVERANS: 'Åtagande som fullgörs under införandet.',
    RUTIN: 'Leverantörens arbetssätt, dokumenterat i docs/rutiner/.',
}


def kravrader(workbook) -> list[dict]:
    """Läser samtliga kravrader ur bilagan, i den ordning de står."""
    rows = []
    for name in SHEETS:
        sheet = workbook[name]
        for row in sheet.iter_rows(min_row=1):
            typ = row[COL_TYPE - 1].value
            text = row[COL_TEXT - 1].value
            if typ not in ('Ska', 'Bör') or not text:
                continue
            rows.append(
                {
                    'sheet': name,
                    'row': row[0].row,
                    'id': (row[COL_ID - 1].value or '(onumrerad)').strip(),
                    'text': ' '.join(str(text).split()),
                    'typ': typ,
                }
            )
    return rows


def main() -> int:
    source = Path(sys.argv[1])
    target = Path(sys.argv[2])

    workbook = openpyxl.load_workbook(source)
    rows = kravrader(workbook)

    saknade = [r['id'] for r in rows if r['id'] not in SVAR]
    if saknade:
        print('Krav utan svar:', ', '.join(saknade))
        return 1

    for entry in rows:
        kategori, kommentar = SVAR[entry['id']]
        sheet = workbook[entry['sheet']]
        sheet.cell(row=entry['row'], column=COL_ANSWER).value = 'Ja'
        sheet.cell(row=entry['row'], column=COL_COMMENT).value = f'{kategori}. {kommentar}'

    workbook.save(target)

    fordelning = Counter(SVAR[r['id']][0] for r in rows)
    skallkrav = sum(1 for r in rows if r['typ'] == 'Ska')
    borkrav = sum(1 for r in rows if r['typ'] == 'Bör')

    print(f'Skrev {len(rows)} svar till {target}')
    print(f'  Ska-krav: {skallkrav}, Bör-krav: {borkrav}')
    for kategori in ORDNING:
        print(f'  {kategori}: {fordelning[kategori]}')

    skriv_matris(rows, fordelning, skallkrav, borkrav)
    return 0


def skriv_matris(rows, fordelning, skallkrav, borkrav) -> None:
    """Skriver docs/kravuppfyllnad.md ur samma källa som bilagan."""
    lines = [
        '# Kravuppfyllnad',
        '',
        'Spårning mot *Bilaga 3. Kravspecifikation*, rad för rad. Samtliga',
        f'{len(rows)} krav i bilagan finns med: {skallkrav} ska-krav och {borkrav} bör-krav.',
        '',
        'Bilagans svarskolumn tillåter bara **Ja** eller **Nej**. Den formen döljer',
        'skillnaden mellan ett krav som är byggt och prövat, ett som vilar på en rutin,',
        'och ett som väntar på uppgifter eller avtal. Därför bär varje rad här också en',
        'bedömning, och samma bedömning står i bilagans kommentarskolumn.',
        '',
        'Tabellen genereras ur `tools/kravsvar.py` med `python3 tools/fyll-bilaga.py`,',
        'så att bilagan och det här dokumentet inte kan säga olika saker.',
        '',
        '## Sammanställning',
        '',
        '| Bedömning | Antal | Innebörd |',
        '| --- | --- | --- |',
    ]
    for kategori in ORDNING:
        lines.append(f'| {kategori} | {fordelning[kategori]} | {FORKLARING[kategori]} |')

    lines += [
        '',
        'Samtliga krav besvaras med **Ja** i bilagan. De rader som är märkta',
        f'*{KONFIGURATION}*, *{AVTAL}* eller *{LEVERANS}* är åtaganden som fullgörs under',
        'införandet, och kommentaren säger i varje enskilt fall exakt vad som återstår.',
        '',
        '## Vad som inte är gjort',
        '',
        'Utan omskrivningar, det här återstår:',
        '',
        '- **Penetrationstest av extern part.** Inte genomfört. Ska beställas inför driftsättning.',
        '- **Oberoende tillgänglighetsgranskning.** Den automatiska granskningen ger noll fel över 33 vyer, men fångar bara ungefär en tredjedel av WCAG.',
        '- **Nativa appar för App Store och Google Play.** Webbappen kan installeras på hemskärmen, men butikspaketeringen är inte gjord.',
        '- **Driftplattform, backuplagring, skanningstjänst, e-post och SMS.** Inte upphandlade.',
        '- **Avtal för BankID, betallösning, digital signering och mätvärden.** Saknas, och de funktionerna är därför avstängda i produkten.',
        '- **Anslutning till Vitec och Aptus.** Adaptern mot Vitec finns med dokumenterad fältavbildning; Aptus-adaptern byggs vid införandet. Båda kräver adress, nyckel och avtal.',
        '',
        '## Krav rad för rad',
        '',
    ]

    nuvarande = None
    for entry in rows:
        avsnitt = entry['sheet']
        if avsnitt != nuvarande:
            nuvarande = avsnitt
            lines += ['', f'### {avsnitt}', '', '| # | Typ | Krav | Bedömning | Var det är löst |', '| --- | --- | --- | --- | --- |']
        kategori, kommentar = SVAR[entry['id']]
        text = entry['text']
        if len(text) > 150:
            text = text[:147] + '…'
        text = text.replace('|', '/')
        lines.append(f"| {entry['id']} | {entry['typ']} | {text} | **{kategori}** | {kommentar} |")

    lines.append('')
    Path('docs/kravuppfyllnad.md').write_text('\n'.join(lines), encoding='utf-8')
    print('Skrev docs/kravuppfyllnad.md')


if __name__ == '__main__':
    raise SystemExit(main())
